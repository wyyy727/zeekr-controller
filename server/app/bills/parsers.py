"""支付宝 / 微信账单解析器。

两个平台的导出格式差异很大，且官方会不定期调整字段。
这里做了**宽容解析**：自动识别编码、自动定位表头行、按列名模糊匹配。

## 支付宝账单特点
- CSV 编码通常是 **GBK**
- 前若干行是说明，表头行以「交易时间」开头
- 金额列可能带正负号，支出为负
- 有「交易对方」「商品说明」「收/支」等列

## 微信账单特点
- CSV 编码通常是 **UTF-8**
- 表头行以「交易时间」开头，前有 16 行左右说明
- 金额列形如 `¥123.45`，需去掉货币符号
- 有「交易对方」「商品」「收/支」等列
"""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Any

# 候选编码，按尝试顺序
CANDIDATE_ENCODINGS = ("utf-8-sig", "utf-8", "gbk", "gb18030", "utf-16")

# 表尾统计行的判定：如「共 123 笔记录」「共12笔」。
#
# 必须锚定「共」+ 数字 + 「笔」，且整行不含其他有意义的交易字段。
# 历史缺陷：只判断「共」与「笔」是否同时出现，会把商户名/备注里恰好
# 含这两个字的**真实交易**一并丢掉（如「共笔文具」「共2笔运费」）。
SUMMARY_ROW_PATTERN = re.compile(r"共\s*\d+\s*笔")

# 列名别名表：统一字段 → 可能的原始列名
COLUMN_ALIASES: dict[str, list[str]] = {
    "timestamp": ["交易时间", "交易创建时间", "支付时间", "时间", "交易日期"],
    "amount": ["金额", "金额(元)", "交易金额", "金额（元）", "发生金额"],
    "direction": ["收/支", "收支", "收/支类型", "资金流向"],
    "merchant": ["交易对方", "对方", "商户名称", "交易对象", "收/付款方"],
    "description": ["商品说明", "商品", "交易说明", "备注", "商品名称", "交易类型"],
    "status": ["交易状态", "当前状态", "状态"],
    "orderNo": ["交易单号", "交易号", "商户订单号", "订单号", "流水号"],
    "channel": ["支付方式", "付款方式", "支付渠道"],
}


def _decode(raw: bytes) -> str | None:
    """自动探测编码并解码；无法可靠解码时返回 None。

    判定分两步，从可靠到宽松：

    1. **严格解码成功**即采纳 —— 这是最可靠的信号（严格模式不会产生替换
       字符，所以旧实现里"替换字符 < 1%"的判定实际等价于"严格成功"）。
    2. 全部严格解码都失败时，退一步用 `errors="replace"` 试一遍，取替换
       字符最少的候选（容忍个别脏字节，但要求脏字符占比 < 1%）。
    3. 连这一步都过不了就返回 None，由调用方给出明确的"编码无法识别"。

    为什么要去掉原来的 `errors="ignore"` 兜底：静默丢弃无法解码的字节，会把
    **编码问题伪装成「未找到表头行」** —— 用户看到的是"格式不对"，而真正的
    原因是文件编码读不了（GBK/UTF-8 混编的账单正是这一类），排障方向被带偏。
    """
    for encoding in CANDIDATE_ENCODINGS:
        try:
            text = raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
        return text or None

    candidates: list[tuple[int, str]] = []
    for encoding in CANDIDATE_ENCODINGS:
        try:
            text = raw.decode(encoding, errors="replace")
        except (UnicodeDecodeError, LookupError):
            continue
        if not text:
            continue
        bad = text.count("\ufffd")
        if bad < len(text) * 0.01:
            candidates.append((bad, text))

    if not candidates:
        return None
    return min(candidates, key=lambda item: item[0])[1]


def _read_xlsx_rows(raw: bytes) -> list[list[str]] | None:
    """用 openpyxl 读取 xlsx，返回与 `csv.reader` 同构的行列表。

    为什么要支持 xlsx：微信导出的「用于个人对账」账单**常见 Excel 格式**，
    而 iOS 的文件选择器也允许选 spreadsheet —— 此前服务端只收 csv/txt，
    客户端能选中却被服务端拒绝，用户会撞上「仅支持 CSV/TXT」的错误。
    两端对齐后，选 xlsx 也能直接导入。

    依赖缺失或文件损坏时返回 None（由调用方给出可读错误），不抛异常。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - 未装 openpyxl 时降级
        return None

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - 非 xlsx / 文件损坏 / 版本不支持
        return None

    try:
        sheet = workbook.active
        if sheet is None:
            return None
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            # 统一成字符串，复用与 CSV 完全相同的后续解析逻辑。
            # 日期单元格经 data_only 得到 datetime，str() 形如
            # "2026-09-01 10:30:00"，正好落在 _parse_datetime 支持的格式里。
            rows.append(["" if value is None else str(value) for value in row])
        return rows
    finally:
        workbook.close()


def _locate_header(rows: list[list[str]]) -> int | None:
    """定位表头行。

    支付宝 / 微信的账单文件前面都有说明文字，
    表头行的特征是**同时包含「交易时间」和时间/金额类列名**。
    """
    for index, row in enumerate(rows[:40]):
        joined = " ".join(cell.strip() for cell in row)
        if "交易时间" in joined and ("金额" in joined or "收/支" in joined or "收支" in joined):
            return index
    # 退而求其次：找任意含「交易时间」的行
    for index, row in enumerate(rows[:40]):
        if any("交易时间" in cell for cell in row):
            return index
    return None


def _map_columns(header: list[str]) -> dict[str, int]:
    """把表头映射为统一字段 → 列索引。"""
    mapping: dict[str, int] = {}
    cleaned = [cell.strip().lstrip("\ufeff") for cell in header]

    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            for index, cell in enumerate(cleaned):
                if cell == alias:
                    mapping[field] = index
                    break
            if field in mapping:
                break

        # 精确匹配失败则退化为包含匹配。
        #
        # 注意不能"第一个命中就用"：表头若没有精确的「金额」而存在
        # 「优惠金额」「退款金额」，会把它们当成交易金额。同一别名命中多列时
        # 选**列名最短**的那个 —— 规范列名通常最短（「金额」<「优惠金额」），
        # 歧义最小。别名之间仍保持原有的优先级顺序。
        if field not in mapping:
            for alias in aliases:
                hits = [
                    index
                    for index, cell in enumerate(cleaned)
                    if cell and alias in cell
                ]
                if hits:
                    mapping[field] = min(hits, key=lambda index: len(cleaned[index]))
                    break

    return mapping


def _parse_amount(text: str) -> float | None:
    """解析金额，去除货币符号与千分位。"""
    if not text:
        return None
    cleaned = re.sub(r"[¥￥,\s元]", "", str(text).strip())
    # 处理括号负数：(123.45)
    negative = cleaned.startswith("(") and cleaned.endswith(")")
    cleaned = cleaned.strip("()")
    try:
        value = float(cleaned)
    except ValueError:
        return None
    return -value if negative else value


def _parse_datetime(text: str) -> datetime | None:
    """解析交易时间。"""
    if not text:
        return None
    text = str(text).strip()
    formats = (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y年%m月%d日 %H:%M:%S",
        "%Y年%m月%d日",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _detect_channel(text: str, source: str) -> str:
    """识别支付渠道。"""
    text = (text or "").lower()
    if "支付宝" in text or source == "alipay":
        return "alipay"
    if "微信" in text or "零钱" in text or source == "wechat":
        return "wechat"
    if "极氪" in text or "zeekr" in text:
        return "zeekr"
    return "other"


def parse_bill(raw: bytes, source: str = "auto") -> dict[str, Any]:
    """解析账单文件。

    Args:
        raw: 文件原始字节
        source: "alipay" / "wechat" / "auto"

    Returns:
        {
            "success": bool,
            "message": str,
            "records": [...],   # 原始解析结果（未过滤充电消费）
            "total": int,
        }
    """
    # xlsx 是 zip 容器，魔数为 PK\x03\x04。先按魔数分流，别把二进制当文本
    # 解码 —— 否则各种编码都会"成功"解出一堆乱码，错误信息完全指错方向。
    if raw[:4] == b"PK\x03\x04":
        rows = _read_xlsx_rows(raw)
        if rows is None:
            return {
                "success": False,
                "message": "无法解析该 Excel 文件，请确认是账单导出文件，"
                           "或用 Excel 另存为「CSV UTF-8」后重试",
                "records": [],
                "total": 0,
            }
    else:
        text = _decode(raw)
        if text is None:
            return {
                "success": False,
                "message": "无法识别文件编码，请用 Excel 打开后另存为「CSV UTF-8」再导入",
                "records": [],
                "total": 0,
            }
        if not text.strip():
            return {"success": False, "message": "文件内容为空", "records": [], "total": 0}

        reader = csv.reader(io.StringIO(text))
        rows = [row for row in reader]

    header_index = _locate_header(rows)
    if header_index is None:
        return {
            "success": False,
            "message": "未找到表头行，请确认是支付宝/微信导出的账单 CSV",
            "records": [],
            "total": 0,
        }

    header = rows[header_index]
    mapping = _map_columns(header)

    if "timestamp" not in mapping or "amount" not in mapping:
        missing = []
        if "timestamp" not in mapping:
            missing.append("交易时间")
        if "amount" not in mapping:
            missing.append("金额")
        return {
            "success": False,
            "message": f"账单缺少必要列：{'、'.join(missing)}",
            "records": [],
            "total": 0,
        }

    def cell(row: list[str], field: str) -> str:
        index = mapping.get(field)
        if index is None or index >= len(row):
            return ""
        return row[index].strip()

    records: list[dict[str, Any]] = []
    for row in rows[header_index + 1:]:
        if not row or not any(cell.strip() for cell in row):
            continue

        timestamp = _parse_datetime(cell(row, "timestamp"))
        amount = _parse_amount(cell(row, "amount"))

        # 跳过表尾统计行（如「共 123 笔记录」）。
        #
        # 用「共 + 数字 + 笔」精确匹配，而非只看到「共」「笔」就丢 —— 后者
        # 会误杀商户名/备注里带这两个字的**真实交易**。再加一道保险：真实
        # 交易必有可解析的金额，统计行没有；两者同时满足才判为统计行。
        if SUMMARY_ROW_PATTERN.search(" ".join(row)) and amount is None:
            continue

        if timestamp is None or amount is None:
            continue

        direction = cell(row, "direction")
        # 只保留支出；微信的「/」表示不计收支，需剔除
        if direction and direction not in ("支出", "支"):
            if "收" in direction and "支" not in direction:
                continue
            if direction in ("/", "不计收支", ""):
                continue

        # 支出取绝对值
        amount = abs(amount)

        records.append({
            "timestamp": timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            "amount": round(amount, 2),
            "merchant": cell(row, "merchant"),
            "description": cell(row, "description"),
            "orderNo": cell(row, "orderNo"),
            "status": cell(row, "status"),
            "channel": _detect_channel(cell(row, "channel"), source),
        })

    return {
        "success": True,
        "message": f"解析成功，共 {len(records)} 笔支出记录",
        "records": records,
        "total": len(records),
    }


def parse_bill_file(path: str, source: str = "auto") -> dict[str, Any]:
    """从文件路径解析账单。"""
    try:
        with open(path, "rb") as handle:
            return parse_bill(handle.read(), source)
    except OSError as exc:
        return {"success": False, "message": f"读取文件失败：{exc}", "records": [], "total": 0}
